"""Redaction: copies a folder of real files into the raw/ of an SSOT instance (or any folder to hand to an
LLM), replacing names, emails and phone numbers in the original formats, and verifies the result.

The originals are never modified. The configuration and the files written next to it (candidates, fake
values, report) contain the real names: keep them outside any repository and outside the destination.
See docs/guide.md for the configuration and the workflow; the command is `ssot-redact` (cli.py).
"""
from __future__ import annotations

import hashlib
import json
import mimetypes
import os
import re
import shutil
import tempfile
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import yaml

# the readers of the files are the ones of ssot-curator: what is verified here is what the curator reads
from curator.config import SourceRule, glob_match
from curator.extractors import extract
from curator.extractors.email import MAIL_SUFFIXES, Message, read_messages
from curator.extractors.text import decode
from curator.extractors.whatsapp import looks_like_chat, parse_chat
from curator.util import fix_ligatures, normalize_for_match

EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")      # without the full stop that may follow
PHONE = re.compile(r"(?<!\d)(?:\+39[\s.]?)?(?:3\d{2}|0\d{1,3})[\s./-]?\d{3,4}[\s./-]?\d{3,4}(?!\d)")
INTL_PHONE = re.compile(r"(?<![\w+])\+\d{1,3}[\s./-]?\d[\d\s./-]{5,14}\d(?!\d)")
# a number is replaced (or covered) as a phone only after one of these words or with a +prefix: codes and
# hexadecimal values of manuals and traces look like phone numbers too
PHONE_CUE = re.compile(r"(?i)\b(?:tel|telefono|phone|cell|cellulare|mobile|mob|fax|whatsapp)\b\.?\s*[:.]?\s*(?:n\.?\s*)?$")


def is_phone(number: str, before: str) -> bool:
    return number.startswith("+") or bool(PHONE_CUE.search(before))


CAPWORD = re.compile(r"\b[A-ZÀ-Ý][a-zà-ÿ]{2,}\b")
UPPERWORD = re.compile(r"\b[A-ZÀ-Ý]{4,}\b")
NAMESEQ = re.compile(r"\b[A-ZÀ-Ý][a-zà-ÿ]+(?:[ \t]+[A-ZÀ-Ý][a-zà-ÿ]+){1,2}\b")
LOWERWORD = re.compile(r"\b[a-zà-ÿ]{3,}\b")
FAKE_DOMAIN = "example.test"


@dataclass
class Sub:
    term: str
    repl: str

    @property
    def rx(self) -> re.Pattern:
        return re.compile(r"(?<![\w])" + re.escape(self.term).replace(r"\ ", r"\s+") + r"(?![\w])", re.IGNORECASE)


@dataclass
class FileResult:
    dest: str
    done: Counter = field(default_factory=Counter)     # text replacements (checked as readable afterwards)
    covered: Counter = field(default_factory=Counter)  # black boxes on images, by replacement
    auto: Counter = field(default_factory=Counter)     # emails and phone numbers replaced automatically
    doubtful: list[str] = field(default_factory=list)  # words read with low OCR confidence
    notes: list[str] = field(default_factory=list)


class Pseudonyms:
    """Emails and phone numbers replaced automatically, each real value always by the same fake one, also
    across runs: the table is kept next to the configuration (and is as sensitive as it)."""

    def __init__(self, path: Path | None, emails: bool = True, phones: bool = True):
        self.path, self.emails, self.phones = path, emails, phones
        data = yaml.safe_load(path.read_text(encoding="utf-8")) if path and path.exists() else None
        data = data or {}
        self.table = {"emails": dict(data.get("emails") or {}), "phones": dict(data.get("phones") or {})}

    def _fake(self, kind: str, value: str, fmt: str) -> str:
        t = self.table[kind]
        if value not in t:
            t[value] = fmt.format(n=len(t) + 1)
        return t[value]

    def apply(self, text: str, res: FileResult) -> str:
        def email(m):
            if m.group(0).lower().endswith("@" + FAKE_DOMAIN):
                return m.group(0)
            res.auto["email"] += 1
            return self._fake("emails", m.group(0).lower(), "person-{n}@" + FAKE_DOMAIN)

        def phone(m):
            if not is_phone(m.group(0), m.string[max(0, m.start() - 25):m.start()]):
                return m.group(0)
            res.auto["phone"] += 1
            return self._fake("phones", re.sub(r"\D", "", m.group(0))[-10:], "phone-{n}")

        if self.emails:
            text = EMAIL.sub(email, text)
        if self.phones:
            text = PHONE.sub(phone, INTL_PHONE.sub(phone, text))
        return text

    def save(self) -> None:
        if self.path and (self.table["emails"] or self.table["phones"]):
            self.path.write_text("# Real value → fake value used in raw/. Keep it outside git: it contains the "
                                 "real emails and phone numbers.\n"
                                 + yaml.safe_dump(self.table, allow_unicode=True, sort_keys=False), encoding="utf-8")


AUTO: Pseudonyms | None = None    # set by run() from the configuration


def apply_text(text: str, subs: list[Sub], res: FileResult) -> str:
    if AUTO is not None:             # first: a name inside an address must not break the address
        text = AUTO.apply(text, res)
    for s in subs:
        text, n = s.rx.subn(s.repl, text)
        if n:
            res.done[s.repl] += n
    return text


# --- per format ---------------------------------------------------------------------------

def redact_text(src: Path, dst: Path, subs, res):
    text, enc = decode(src.read_bytes())
    out = apply_text(text, subs, res)
    enc = "utf-8" if enc == "utf-8-sig" else enc
    dst.write_bytes(out.encode(enc, errors="strict"))


def _docx_paragraphs(doc):
    def from_container(c):
        for p in c.paragraphs:
            yield p
        for t in c.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from from_container(cell)
    yield from from_container(doc)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            yield from from_container(part)


def redact_docx(src: Path, dst: Path, subs, res):
    import docx
    d = docx.Document(str(src))
    for p in _docx_paragraphs(d):
        new = apply_text(p.text, subs, FileResult(""))
        if new != p.text:
            apply_text(p.text, subs, res)        # counting
            runs = p.runs
            if runs:
                runs[0].text = new
                for r in runs[1:]:
                    r.text = ""
    cp = d.core_properties
    for attr in ("author", "last_modified_by", "comments", "keywords", "subject", "category", "identifier"):
        setattr(cp, attr, "")
    cp.title = apply_text(cp.title or "", subs, res)
    d.save(str(dst))


def redact_xlsx(src: Path, dst: Path, subs, res):
    import openpyxl
    wb = openpyxl.load_workbook(str(src))
    for ws in wb.worksheets:
        new_title = apply_text(ws.title, subs, res)
        if new_title != ws.title:
            ws.title = new_title
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    c.value = apply_text(c.value, subs, res)
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.save(str(dst))


_PDF_CHARS = str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-",
                            "…": "...", " ": " ", "•": "-"})


def redact_pdf(src: Path, dst: Path, subs, res):
    from pypdf import PdfReader
    from ssot_redact.pdfwrite import text_pdf
    reader = PdfReader(str(src))
    pages = []
    for pg in reader.pages:
        text = apply_text(fix_ligatures(pg.extract_text() or ""), subs, res).translate(_PDF_CHARS)
        pages.append((text, float(pg.mediabox.width), float(pg.mediabox.height)))
    dst.write_bytes(text_pdf(pages))          # no metadata: written by hand
    res.notes.append(f"PDF rebuilt as text only ({len(reader.pages)} pages, layout and images dropped)")


# --- mail: rebuilt with only what the curator reads ------------------------------------------------

def _fake_id(value: str) -> str:
    v = value.strip().strip("<>").lower()
    return f"<{hashlib.sha256(v.encode('utf-8')).hexdigest()[:16]}@{FAKE_DOMAIN}>" if v else ""


def _redacted_message(m: Message, subs, res, tmp: Path):
    from email.message import EmailMessage
    from email.utils import format_datetime
    from email.utils import formataddr, getaddresses
    out = EmailMessage()
    for header, value in (("From", m.sender), ("To", m.to), ("Cc", m.cc)):
        if value:              # address by address and quoted: a replacement like "[OMITTED]" must not break it
            out[header] = ", ".join(formataddr((apply_text(n, subs, res), apply_text(a, subs, res)))
                                    for n, a in getaddresses([value]) if n or a)
    if m.subject:
        out["Subject"] = apply_text(m.subject, subs, res)
    if m.date:
        out["Date"] = format_datetime(m.date)
    for header, value in (("Message-ID", _fake_id(m.message_id)), ("In-Reply-To", _fake_id(m.in_reply_to)),
                          ("References", " ".join(_fake_id(r) for r in m.references))):
        if value:
            out[header] = value
    out.set_content(apply_text(m.body, subs, res))
    for i, a in enumerate(m.attachments):
        name = apply_text(a.name, subs, res)
        fn = REDACTORS.get(Path(a.name).suffix.lower())
        if fn is None or fn is redact_mail:
            res.notes.append(f"attachment dropped (format not handled): {name}")
            continue
        src, dst = tmp / f"in{i}{Path(a.name).suffix}", tmp / f"out{i}{Path(a.name).suffix}"
        src.write_bytes(a.data)
        try:
            fn(src, dst, subs, res)
        except Exception as e:  # one damaged attachment must not stop the mail
            res.notes.append(f"attachment dropped (unreadable: {e}): {name}")
            continue
        maintype, _, subtype = (mimetypes.guess_type(name)[0] or "application/octet-stream").partition("/")
        out.add_attachment(dst.read_bytes(), maintype=maintype, subtype=subtype, filename=name)
    return out


def redact_mail(src: Path, dst: Path, subs, res):
    import mailbox
    msgs = read_messages(src)
    with tempfile.TemporaryDirectory() as tmp:
        out = [_redacted_message(m, subs, res, Path(tmp)) for m in msgs]
    if src.suffix.lower() == ".mbox":
        box = mailbox.mbox(str(dst))
        for m in out:
            box.add(m)
        box.flush()
        box.close()
    else:
        dst.write_bytes(out[0].as_bytes())
    res.notes.append(f"mail rebuilt ({len(msgs)} message(s)): other headers and formatting dropped")


# --- images: OCR + black boxes ----------------------------------------------------------------

OCR = None            # TesseractOcr set by run() from the configuration


class NeedsOcr(Exception):
    """An image or a scanned PDF without `ocr:` in the configuration: it is skipped, not copied as it is."""
LOW_CONF = 60.0


def _tok(s: str) -> str:
    return re.sub(r"[^\w]", "", normalize_for_match(s))


def boxes_to_cover(words, subs: list[Sub], res: FileResult) -> list[tuple[int, int, int, int]]:
    """Boxes (left, top, right, bottom) of the words to cover: substituted terms (also over several
    consecutive words of a line), emails and phone numbers."""
    by_line: dict[tuple, list] = {}
    for w in words:
        by_line.setdefault(w.line, []).append(w)
    out = []
    for line in by_line.values():
        toks = [_tok(w.text) for w in line]
        for s in subs:
            want = [t for t in (_tok(x) for x in s.term.split()) if t]
            for i in range(len(toks) - len(want) + 1):
                if want and toks[i:i + len(want)] == want:
                    ws = line[i:i + len(want)]
                    out.append((min(w.left for w in ws), min(w.top for w in ws),
                                max(w.left + w.width for w in ws), max(w.top + w.height for w in ws)))
                    res.covered[s.repl] += 1
        for j, w in enumerate(line):
            before = " ".join(x.text for x in line[max(0, j - 2):j]) + " "
            if EMAIL.search(w.text) or ((PHONE.search(w.text) or INTL_PHONE.search(w.text))
                                        and is_phone(w.text, before)):
                out.append((w.left, w.top, w.left + w.width, w.top + w.height))
                res.covered["[email/phone]"] += 1
            if w.conf < LOW_CONF and len(_tok(w.text)) >= 3:
                res.doubtful.append(w.text)
    return out


def _redacted_image(source, subs, res):
    """The image (path or bytes) with black boxes on the words to redact."""
    import io
    from PIL import Image, ImageDraw
    im = Image.open(io.BytesIO(source) if isinstance(source, bytes) else source).convert("RGB")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    d = ImageDraw.Draw(im)
    for l, t, r, b in boxes_to_cover(OCR.read(buf.getvalue(), suffix=".png").words, subs, res):
        d.rectangle((l - 3, t - 3, r + 3, b + 3), fill="black")
    return im


def redact_image(src: Path, dst: Path, subs, res):
    if OCR is None:
        raise NeedsOcr("image")
    im = _redacted_image(src, subs, res)
    fmt = "JPEG" if src.suffix.lower() in (".jpg", ".jpeg") else "PNG"
    im.save(dst, fmt, **({"quality": 95} if fmt == "JPEG" else {}))     # saved without EXIF/metadata
    res.notes.append("image: OCR + black boxes, metadata dropped")


def redact_scanned_pdf(src: Path, dst: Path, subs, res):
    """A PDF without text (scan): each page image is OCR'd, covered, and the PDF rebuilt from them."""
    import io
    import pypdfium2
    if OCR is None:
        raise NeedsOcr("scanned PDF")
    pages = []
    doc = pypdfium2.PdfDocument(str(src))
    try:
        for pg in doc:        # rendered: the scan may sit in nested objects that a plain extraction misses
            buf = io.BytesIO()
            pg.render(scale=200 / 72).to_pil().save(buf, "PNG")
            pages.append(_redacted_image(buf.getvalue(), subs, res))
    finally:
        doc.close()
    pages[0].save(dst, "PDF", save_all=True, append_images=pages[1:], resolution=150)
    res.notes.append(f"scanned PDF: OCR + black boxes on {len(pages)} pages, rebuilt from the images")


def redact_any_pdf(src: Path, dst: Path, subs, res):
    from pypdf import PdfReader
    has_text = any((p.extract_text() or "").strip() for p in PdfReader(str(src)).pages)
    (redact_pdf if has_text else redact_scanned_pdf)(src, dst, subs, res)


REDACTORS = {".txt": redact_text, ".csv": redact_text, ".md": redact_text,
             ".docx": redact_docx, ".xlsx": redact_xlsx, ".pdf": redact_any_pdf,
             ".png": redact_image, ".jpg": redact_image, ".jpeg": redact_image,
             ".eml": redact_mail, ".mbox": redact_mail, ".msg": redact_mail}
IMAGES = (".png", ".jpg", ".jpeg")


# --- which files ------------------------------------------------------------------------------

def map_path(rel: str, rules: list[dict]) -> str:
    for r in rules:
        f = r["from"].strip("/")
        if rel.lower() == f.lower() or rel.lower().startswith(f.lower() + "/"):
            return (r["to"].strip("/") + rel[len(f):]).lstrip("/")
    return rel


def select(cfg: dict, src_root: Path, subs: list[Sub]) -> list[tuple[str, str]]:
    """(path in the sample, path in raw/): the listed `files`, or every file by include/exclude/map. The
    names in the destination path are redacted; a .msg becomes .eml."""
    if "files" in cfg:
        pairs = [(i["from"], i["to"]) for i in cfg["files"]]
    else:
        include, exclude = cfg.get("include") or ["**/*"], cfg.get("exclude") or []
        pairs = []
        for p in sorted(src_root.rglob("*")):
            rel = p.relative_to(src_root).as_posix()
            if (not p.is_file() or p.name.startswith("~$")              # Office lock files
                    or not any(glob_match(x, rel) for x in include) or any(glob_match(x, rel) for x in exclude)):
                continue
            pairs.append((rel, map_path(rel, cfg.get("map") or [])))
    out = []
    for src, dst in pairs:
        dst = "/".join(apply_text(part, subs, FileResult("")) for part in dst.split("/"))
        if dst.lower().endswith(".msg"):
            dst = dst[:-4] + ".eml"
        out.append((src, dst))
    return out


# --- discovery: the names to redact ----------------------------------------------------------------

def discover(cfg_path: Path, cfg: dict, pairs: list[tuple[str, str]], src_root: Path, subs: list[Sub]) -> Path:
    """Candidates for `substitutions:` from the whole sample, most widespread first; nothing is written to raw/."""
    from email.utils import getaddresses
    people: dict[str, set] = defaultdict(set)          # name → how it was found
    people_files: dict[str, set] = defaultdict(set)
    seq_files: dict[str, set] = defaultdict(set)
    word_files: dict[str, set] = defaultdict(set)
    seqs, words, domains, lower = Counter(), Counter(), Counter(), Counter()
    read, not_read = 0, Counter()

    def person(name: str, how: str, rel: str):
        name = " ".join(name.replace('"', "").split())
        if len(name) >= 3 and not EMAIL.fullmatch(name):
            people[name].add(how)
            people_files[name].add(rel)

    def address(addr: str, rel: str):
        local, _, domain = addr.lower().partition("@")
        if domain and domain != FAKE_DOMAIN:
            domains[domain] += 1
        bits = [b for b in re.split(r"[._-]", local) if b.isalpha() and len(b) >= 2]
        if len(bits) >= 2:
            person(" ".join(b.capitalize() for b in bits), "mail address", rel)

    for i, (rel, _) in enumerate(pairs, start=1):
        if i % 200 == 0:
            print(f"  read {i}/{len(pairs)}")
        p, ext = src_root / rel, Path(rel).suffix.lower()
        if ext not in REDACTORS or (ext in IMAGES and OCR is None):
            not_read[ext or "(no extension)"] += 1
            continue
        try:
            text = all_text(p)
            if ext in MAIL_SUFFIXES:
                for m in read_messages(p):
                    for value in (m.sender, m.to, m.cc):
                        for name, addr in getaddresses([value]):
                            if name:
                                person(name, "mail header", rel)
                            if addr:
                                address(addr, rel)
            elif ext == ".txt" and looks_like_chat(p):
                for msg in parse_chat(decode(p.read_bytes())[0]):
                    person(msg.author, "chat author", rel)
        except Exception:  # a damaged file is reported by the real run
            not_read["(unreadable)"] += 1
            continue
        read += 1
        for e in EMAIL.findall(text):
            address(e, rel)
        text = EMAIL.sub(" ", text)      # "mario.rossi@…" does not make "rossi" an ordinary word
        lower.update(LOWERWORD.findall(text))
        for s in NAMESEQ.findall(text):
            s = " ".join(s.split())
            seqs[s] += 1
            seq_files[s].add(rel)
        for w in CAPWORD.findall(text) + UPPERWORD.findall(text):
            words[w] += 1
            word_files[w].add(rel)

    allowed = {a.lower() for a in cfg.get("allowed", [])}
    known = {s.term.lower() for s in subs}

    def ordinary(w: str) -> bool:
        """Written in lowercase often enough to be a common word: a name written once in lowercase is not."""
        low, cap = lower[w.lower()], words[w.capitalize()] + words[w.upper()]
        return low >= 2 and low * 4 >= cap

    def fresh(name: str) -> bool:
        return name.lower() not in known and name.lower() not in allowed

    def q(s: str) -> str:
        return json.dumps(s, ensure_ascii=False)

    rank = sorted(people, key=lambda n: (-len(people_files[n]), n))
    lines = [f"# Candidates for redaction: `ssot-redact discover` read {read} files of {cfg_path.name}; nothing",
             "# was written to raw/. Copy the lines to keep into `substitutions:` (longest forms first: a full name",
             "# before its surname), choose the replacements, and put the words that are fine into `allowed:`.",
             "# Emails and phone numbers are replaced automatically (`auto:`). This file contains real names: keep",
             "# it outside git.",
             "", "people:   # mail senders and recipients, mail addresses, chat authors: the most reliable"]
    n = 0
    for name in rank:
        if fresh(name):
            n += 1
            lines.append(f'  - [{q(name)}, "Person {n}"]   # {len(people_files[name])} files; '
                         f'{", ".join(sorted(people[name]))}')
    lines += ["", "names:    # two or three capitalized words in a row, none of them a common lowercase word"]
    for s in sorted(seqs, key=lambda s: (-len(seq_files[s]), -seqs[s], s)):
        if fresh(s) and s not in people and not any(ordinary(w) for w in s.split()):
            n += 1
            lines.append(f'  - [{q(s)}, "Person {n}"]   # {len(seq_files[s])} files, {seqs[s]} times')
    lines += ["", "words:    # single capitalized words (almost) never written in lowercase: names, places, companies or",
              "          # technical terms. The replacement is to choose: \"?\" is not a valid one"]
    for w in sorted(words, key=lambda w: (-len(word_files[w]), -words[w], w)):
        if fresh(w) and not ordinary(w):
            lines.append(f'  # - [{q(w)}, "?"]   # {len(word_files[w])} files, {words[w]} times')
    lines += ["", "domains:  # mail domains: the organizations involved (their names are probably among the words)"]
    for d, c in domains.most_common():
        lines.append(f"  # - [{q(d)}, \"?\"]   # {c} addresses")
    if not_read:
        lines += ["", "# Not read (format not handled, images without `ocr:`, unreadable): "
                  + ", ".join(f"{k} ×{v}" for k, v in not_read.most_common())]
    out = cfg_path.with_name(cfg_path.stem + "-candidates.yaml")
    out.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
    print(f"Files read: {read} | not read: {sum(not_read.values())} | people: {len(rank)} | "
          f"names: {len(seqs)} | words: {len(words)}")
    print(f"Candidates: {out}")
    return out


# --- verification ---------------------------------------------------------------------------

def all_text(path: Path) -> str:
    """All the text the curator (or anyone opening the file) can see, metadata included."""
    ext = path.suffix.lower()
    parts = [path.name]
    if ext in MAIL_SUFFIXES:
        try:
            msgs = read_messages(path)
        except Exception as e:  # verification must not stop on one file
            return f"{path.name}\n[extraction failed: {e}]"
        with tempfile.TemporaryDirectory() as tmp:
            for m in msgs:
                parts += [m.sender, m.to, m.cc, m.subject, m.message_id, m.in_reply_to, " ".join(m.references),
                          m.body]
                for i, a in enumerate(m.attachments):
                    p = Path(tmp) / f"{i}-{a.name}"
                    p.write_bytes(a.data)
                    parts.append(all_text(p))
        return "\n".join(parts)
    kind = {".pdf": "pdf", ".docx": "docx", ".xlsx": "table", ".csv": "table", ".png": "image", ".jpg": "image",
            ".jpeg": "image"}.get(ext, "text")
    try:
        parts.append(extract(path, path.name, SourceRule("x", kind, None, True, {}), reader=OCR,
                             reader_mode="ocr" if OCR else "off").text)
    except Exception as e:  # verification must not stop on one file
        parts.append(f"[extraction failed: {e}]")
    try:
        if ext in (".docx", ".xlsx"):
            with zipfile.ZipFile(path) as z:
                for n in z.namelist():
                    if n.endswith((".xml", ".rels")):
                        xml = z.read(n).decode("utf-8", "replace")
                        parts.append(re.sub(r"<[^>]+>", " ", xml))
        if ext == ".pdf":
            from pypdf import PdfReader
            r = PdfReader(str(path))
            parts.append(" ".join(str(v) for v in (r.metadata or {}).values() if v))
            xmp = r.trailer["/Root"].get("/Metadata")
            if xmp is not None:
                parts.append(xmp.get_object().get_data().decode("utf-8", "replace"))
    except Exception as e:
        parts.append(f"[metadata unreadable: {e}]")
    return "\n".join(parts)


def verify(dest_root: Path, forbidden: list[str], allowed: set[str], results: list[FileResult] = ()) -> dict:
    """Forbidden terms still readable in raw/, emails, phones and capitalized words to review, and the
    replacements made that cannot be read back (text lost). Each file is read once."""
    leaks, lost, emails, phones, words = [], [], Counter(), Counter(), Counter()
    by_norm = {normalize_for_match(f): f for f in forbidden if normalize_for_match(f)}
    one_rx = re.compile(r"(?<!\w)(" + "|".join(re.escape(n) for n in sorted(by_norm, key=len, reverse=True))
                        + r")(?!\w)") if by_norm else None
    # the space-less comparison (for "J O H N") only for long terms: "ana" is inside "banana"
    squashed_terms = [(re.sub(r"\s+", "", n), f) for n, f in by_norm.items() if len(re.sub(r"\s+", "", n)) >= 6]
    expected = {r.dest: r.done for r in results}
    files = sorted(p for p in dest_root.rglob("*") if p.is_file())
    for i, p in enumerate(files, start=1):
        if i % 200 == 0:
            print(f"  verified {i}/{len(files)}")
        rel = p.relative_to(dest_root).as_posix()
        text = all_text(p)
        nt = normalize_for_match(text)
        found = {by_norm[m.group(1)] for m in one_rx.finditer(nt)} if one_rx else set()
        if squashed_terms:
            squashed = re.sub(r"\s+", "", nt)
            found |= {f for sq, f in squashed_terms if sq in squashed}
        leaks += [(rel, f) for f in sorted(found, key=str.lower)]
        for repl, n in expected.get(rel, {}).items():
            if text.count(repl) < n:
                lost.append(f"{rel}: «{repl}» replaced {n} times but readable only {text.count(repl)}")
        for e in EMAIL.findall(text):
            if not e.lower().endswith("@" + FAKE_DOMAIN):
                emails[e] += 1
        for t in PHONE.findall(text):
            phones[t.strip()] += 1
        for w in CAPWORD.findall(text):
            if w not in allowed:
                words[w] += 1
    return {"leaks": leaks, "lost": lost, "emails": emails, "phones": phones, "words": words}


# --- main -------------------------------------------------------------------------------------

MODES = ("discover", "run", "add", "verify")


def run(config: str | Path, mode: str = "run") -> int:
    """`mode`: discover (candidates only), run (copy + redact + verify), add (only the files not yet in the
    destination), verify. Exit code: 0 ok, 1 refused, 2 forbidden terms or lost replacements found."""
    if mode not in MODES:
        raise ValueError(f"unknown mode: {mode}")
    cfg_path = Path(config).resolve()
    only_verify, add, find = mode == "verify", mode == "add", mode == "discover"
    cfg = yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
    global OCR, AUTO
    OCR = None
    if cfg.get("ocr"):
        from curator.ocr import TesseractOcr
        OCR = TesseractOcr(cfg["ocr"].get("command", "tesseract"), cfg["ocr"].get("languages", "eng"))
    auto = cfg.get("auto") or {}
    AUTO = Pseudonyms(cfg_path.with_name(cfg_path.stem + "-pseudonyms.yaml"),
                      auto.get("emails", True), auto.get("phones", True))
    base = cfg_path.parent
    src_root = (base / cfg["sample"]).resolve()
    dest_root = (base / cfg["destination"]).resolve()
    subs = [Sub(a, b) for a, b in cfg.get("substitutions", [])]

    if find:
        AUTO = None                    # names are looked for in the original text
        discover(cfg_path, cfg, select(cfg, src_root, []), src_root, subs)
        return 0

    results: list[FileResult] = []
    skipped: dict[str, list[str]] = defaultdict(list)     # extension → files not copied
    failed: list[str] = []
    if not only_verify:
        if not add and dest_root.exists() and any(p.is_file() for p in dest_root.rglob("*")):
            print(f"ERROR: {dest_root} already contains files. Empty it by hand to regenerate it, "
                  "or use `ssot-redact add` for the files not yet there.")
            return 1
        pairs = select(cfg, src_root, subs)
        copy_unknown = cfg.get("unknown_formats", "skip") == "copy"
        done_by_hash: dict[str, tuple[Path, FileResult]] = {}
        for i, (src_rel, dst_rel) in enumerate(pairs, start=1):
            if i % 50 == 0:
                print(f"  {i}/{len(pairs)}")
            src = src_root / src_rel
            dst = dest_root / dst_rel
            if add and dst.exists():
                continue
            if not src.is_file():
                print(f"ERROR: file not found in the sample: {src_rel}")
                return 1
            fn = REDACTORS.get(src.suffix.lower())
            if fn is None and not copy_unknown:
                skipped[src.suffix.lower() or "(no extension)"].append(src_rel)
                continue
            dst.parent.mkdir(parents=True, exist_ok=True)
            res = FileResult(dst_rel)
            h = hashlib.sha256(src.read_bytes()).hexdigest()
            if h in done_by_hash:
                # identical copies in the sample stay identical in raw/ (same document for the curator)
                first, first_res = done_by_hash[h]
                shutil.copyfile(first, dst)
                res.done = first_res.done.copy()
                res.notes.append(f"identical copy of {first_res.dest}")
            elif fn is None:
                shutil.copyfile(src, dst)
                res.notes.append("copied WITHOUT redaction (format not handled)")
            else:
                try:
                    fn(src, dst, subs, res)
                except NeedsOcr as e:
                    dst.unlink(missing_ok=True)
                    skipped[f"{e} without `ocr:`"].append(src_rel)
                    continue
                except Exception as e:  # one damaged file must not stop thousands
                    dst.unlink(missing_ok=True)
                    failed.append(f"{src_rel}: {type(e).__name__}: {e}")
                    continue
                done_by_hash[h] = (dst, res)
            st = src.stat()
            os.utime(dst, (st.st_atime, st.st_mtime))   # the curator needs the date of the original file
            results.append(res)
        AUTO.save()

    forbidden = sorted({s.term for s in subs} | set(cfg.get("forbidden_extra", [])) | set(AUTO.table["emails"]),
                       key=str.lower)
    v = verify(dest_root, forbidden, set(cfg.get("allowed", [])), results)
    lost = v["lost"]
    lines = ["# Redaction report", "", f"Destination: `{dest_root}`", ""]
    if results or skipped or failed:
        lines += ["## Unreadable replacement text", ""] + ([f"- {x}" for x in lost] or ["None."]) + [""]
        lines += ["## Failed (not copied)", ""] + ([f"- {x}" for x in failed] or ["None."]) + [""]
        lines += ["## Not copied (format not handled, or images and scans without `ocr:`)", ""]
        for ext, names in sorted(skipped.items(), key=lambda x: (-len(x[1]), x[0])):
            lines += [f"- **{ext}** ×{len(names)}: " + ", ".join(names)]
        lines += (["None."] if not skipped else []) + [""]
        if AUTO.table["emails"] or AUTO.table["phones"]:
            lines += [f"Emails replaced automatically: {len(AUTO.table['emails'])} different, phone numbers: "
                      f"{len(AUTO.table['phones'])} (table: `{AUTO.path.name}`). Codes and serial numbers can "
                      "look like phone numbers: check the table, and if they are not phones set "
                      "`auto: {phones: false}` and redo.", ""]
        lines += ["## Replacements made", "", "| file | replacements |", "|---|---|"]
        for r in results:
            done = ", ".join(f"{k} ×{n}" for k, n in sorted((r.done + r.covered + r.auto).items())) or "none"
            lines.append(f"| {r.dest} | {done}{' — ' + '; '.join(r.notes) if r.notes else ''} |")
        lines.append("")
        doubtful = [r for r in results if r.doubtful]
        if doubtful:
            lines += ["## Images: words read with low OCR confidence (check them on the image)", ""]
            lines += [f"- {r.dest}: " + ", ".join(sorted(set(r.doubtful))) for r in doubtful] + [""]
    lines += ["## Verification: forbidden terms still present", ""]
    lines += [f"- **{rel}**: «{term}»" for rel, term in v["leaks"]] or ["None."]
    lines += ["", "## To review by hand", "", "Emails found:"]
    lines += [f"- {e} ×{n}" for e, n in v["emails"].most_common()] or ["- none"]
    lines += ["", "Possible phone numbers:"]
    lines += [f"- {t} ×{n}" for t, n in v["phones"].most_common()] or ["- none"]
    lines += ["", "Capitalized words (possible names), most frequent first:", ""]
    lines += [", ".join(f"{w} ×{n}" for w, n in sorted(v["words"].items(), key=lambda x: (-x[1], x[0])))]
    report = "\n".join(lines) + "\n"
    out = cfg_path.with_name(cfg_path.stem + "-report.md")
    out.write_text(report, encoding="utf-8", newline="\n")
    print(f"Files prepared: {len(results)} | not copied (format): {sum(len(x) for x in skipped.values())} | "
          f"failed: {len(failed)} | forbidden terms found: {len(v['leaks'])} | "
          f"unreadable replacements: {len(lost)} | "
          f"emails: {sum(v['emails'].values())} | phones: {sum(v['phones'].values())} | "
          f"capitalized words to review: {len(v['words'])}")
    print(f"Report: {out}")
    return 2 if v["leaks"] or lost else 0

